import ast
import os
from copy import deepcopy

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..core.data import Info, InfoComposition, InfoStr
from .data import ComtesData, list_of_tests


def convert_materials(compositions: list[InfoComposition]) -> list[Info]:
    compositions_new = []
    for composition_old in compositions:
        composition_new = {}
        multiple_values = False
        for material, value in composition_old.value:
            # If there are no digits, skip
            if not any(ch.isdigit() for ch in value):
                continue
            # Try to convert the value into a float
            try:
                value = float(value.replace(",", "."))
            except Exception:
                pass
            # Check whether the value is there already
            if material in composition_new:
                if composition_new[material] != value:
                    multiple_values = True
                else:
                    continue
            # Assign the value
            composition_new[material] = value
        # If there are multiple values, ignore it
        if multiple_values:
            print(f"Multiple elements detected for: {composition_old}")
            continue
        # Add the new composition into the list
        if len(composition_new) >= 1:
            compositions_new.append(Info(composition_new, composition_old.source))
    return _merge_adjacent(compositions_new)


def _merge_adjacent(infos: list[Info]) -> list[Info]:
    if not infos:
        return []

    infos_new = [deepcopy(infos[0])]
    for i in range(1, len(infos)):
        new_dict = infos_new[-1].value
        old_dict = infos[i].value
        if infos[i].source == infos_new[-1].source and new_dict.keys().isdisjoint(old_dict.keys()):
            new_dict.update(old_dict)
        else:
            infos_new.append(deepcopy(infos[i]))

    return infos_new


def export_dataframe(df: pd.DataFrame, filename: str, replace_source: list[tuple[str, str]] = []) -> None:
    cols_write = ["Material", "Material EDA", "Composition", "Composition EDA", "Test", "Specimen"]

    # Format the columns
    df_values = df.copy()
    for col in cols_write:
        mask = df_values[col].isnull()
        for i in mask[mask].index:
            df_values.at[i, col] = []
        df_values[col] = df_values[col].apply(lambda xs: [str(x.value) for x in xs])
        df_values[col] = df_values[col].apply(lambda xs: "\n".join(xs))
        df_values[col] = df_values[col].replace({"": " "})

    # Save DataFrame to Excel
    df_values.to_excel(filename, index=False)

    # Open with openpyxl
    wb = load_workbook(filename)
    ws = wb.active
    assert ws is not None

    # Add hyperlinks
    cols = df.columns
    for i, (_, df_row) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(cols, start=1):
            value = df_row[col]
            if isinstance(value, list) and len(value) == 1:
                cell = ws.cell(row=i, column=j)
                source = value[0].source
                for x, y in replace_source:
                    source = source.replace(x, y)
                cell.hyperlink = source
                cell.font = Font(color="0000FF", underline="single")

    # Define alignment: left, top, wrapped (prevents spill)
    text_alignment_no_wrap = Alignment(horizontal="left", vertical="top", wrap_text=False)
    text_alignment_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Apply formatting to all used cells
    baseline_width = 15
    max_len_list = df.map(lambda x: len(x) if isinstance(x, list) else 0).max(axis=1)
    max_len_list = [1] + list(max_len_list.to_numpy())
    for i, (row, n) in enumerate(zip(ws.iter_rows(), max_len_list), start=1):
        ws.row_dimensions[i].height = baseline_width * max(n, 1)
        for cell in row:
            cell.number_format = "@"
            if isinstance(cell.value, str) and "\n" in cell.value:
                cell.alignment = text_alignment_wrap
            else:
                cell.alignment = text_alignment_no_wrap

    # Convert range to Excel Table
    last_row = ws.max_row
    last_col = ws.max_column
    end_col_letter = get_column_letter(last_col)
    ref = f"A1:{end_col_letter}{last_row}"

    table = Table(displayName="ExportedTable", ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    # Alternate the row colours
    color1 = PatternFill(fill_type="solid", fgColor="E8DAEF")
    color2 = PatternFill(fill_type="solid", fgColor="FCF3CF")
    current_fill = color1
    previous_value = None
    for row in range(2, last_row + 1):
        file_value = ws.cell(row=row, column=1).value
        if previous_value is not None and file_value != previous_value:
            current_fill = color2 if current_fill == color1 else color1
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).fill = current_fill
        previous_value = file_value

    wb.save(filename)


def get_conversion_tests() -> dict[str, str]:
    conversion_tests = {}
    for x, y in list_of_tests:
        conversion_tests[x] = y
    return conversion_tests


def _resource_path(folder: str | None, file_name: str) -> str:
    if folder is None:
        folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resources")
    path = os.path.join(folder, file_name)
    if not os.path.exists(path):
        raise FileNotFoundError(f"Resource file {path} not found. The COMTES resources are not part of the repository and must be copied there manually.")
    return path


def load_eda_composition(folder: str | None = None) -> dict[str, dict[str, str]]:
    composition1 = pd.read_excel(_resource_path(folder, "Chem_composition_COMTES_materials.xlsx"))
    composition1["Name"] = composition1["Name"].apply(lambda x: ast.literal_eval(x)["value"])
    composition2 = pd.read_excel(_resource_path(folder, "Chem_composition_STEELS_EDA.xlsx"))
    composition2["Name"] = composition2["Name"].apply(lambda x: ast.literal_eval(x)["value"])
    composition = pd.concat((composition1, composition2))
    composition = composition.drop_duplicates()

    if len(composition) != composition["Name"].nunique():
        raise ValueError("There are multiple values for column Name")

    cols = composition.columns
    cols_skip = ["Name", "Alloy Type"]
    results = {}
    for _, row in composition.iterrows():
        result = {}
        for col in cols:
            if col in cols_skip:
                continue
            if not pd.isnull(row[col]):
                result[col.rstrip(" [%]")] = row[col]
        results[row["Name"]] = result
    return results


def load_material_conversion(path: str | None = None) -> pd.DataFrame:
    if path is None:
        path = _resource_path(None, "conversion.csv")
    return pd.read_csv(path)


def process_files(root: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    file_names = os.listdir(root)
    file_names = [x for x in file_names if x.endswith(".pkl")]
    results = []
    values_tests_all = pd.DataFrame()
    for file_name in file_names:
        order_name = os.path.splitext(file_name)[0]
        data = ComtesData.load(f"{root}/{file_name}", None)
        # Fix integers which should not be there
        for info in data.material:
            info.to_str()
            info.normalize_value()
        tests = []
        for data_file_name, (_, _, test_type) in zip(data.file_names, data.tests):
            if test_type != "":
                tests.append(InfoStr(test_type, data_file_name))
        results.append(
            {
                "File": order_name,
                "Test": tests,
                "Material": data.material,
                "Composition": data.composition,
            }
        )
        values_tests = data.values_tests.copy()
        values_tests = values_tests.loc[:, values_tests.columns.notna() * (values_tests.columns != "nan")]
        values_tests = values_tests.rename({"d0": "aux_D0", "du": "aux_Du", "bu": "aux_Bu"}, axis=1)
        # TODO: write better
        for col in ["n", "C"]:
            if col in values_tests.columns:
                values_tests = values_tests.drop(col, axis=1)
        values_tests.insert(0, "File", order_name)
        values_tests_all = pd.concat((values_tests_all, values_tests))
    values_tests_all["Specimen"] = values_tests_all["Specimen"].apply(lambda x: [x])
    return pd.DataFrame(results), values_tests_all
